from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"F:\UnityGame\Chimera\output\任务二_立项汇报_四人讲稿_口语版.docx")
FONT = "Microsoft YaHei"
BLACK = "000000"
MUTED = "5B6870"
TEAL = "138A8A"
PALE_BLUE = "EAF5F7"
PALE_GRAY = "F6F8F9"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", *, style=None, size=11, bold=False, color=BLACK,
                  before=0, after=6, line=1.35, keep=False, indent=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_cue(doc, text):
    return add_paragraph(doc, text, size=9.5, bold=True, color=TEAL, before=5, after=4, line=1.1, keep=True)


def add_speaker_header(doc, number, title, pages, timing):
    add_paragraph(doc, f"第{number}部分  {title}", style="Heading 1", size=15, bold=True,
                  before=0, after=4, line=1.15, keep=True)
    p = add_paragraph(doc, f"对应PPT第{pages}页    建议时长{timing}    主讲人________",
                      size=9.5, color=MUTED, before=0, after=12, line=1.15, keep=True)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
styles["Normal"].font.name = FONT
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Normal"].font.size = Pt(11)
styles["Title"].font.name = FONT
styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Title"].font.color.rgb = RGBColor.from_string(BLACK)
styles["Title"].font.size = Pt(20)
styles["Title"].font.bold = True
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
styles["Heading 1"].font.name = FONT
styles["Heading 1"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Heading 1"].font.color.rgb = RGBColor.from_string(BLACK)
styles["Heading 1"].font.size = Pt(15)
styles["Heading 1"].font.bold = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
run = title.add_run("任务二科学投喂与鱼群摄食监测立项汇报讲稿")
set_run_font(run, size=20, bold=True, color=BLACK)
direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
if direct_border is not None:
    title._p.get_or_add_pPr().remove(direct_border)

subtitle = add_paragraph(doc, "四人汇报口语版", size=12, color=MUTED, after=16, line=1.1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_paragraph(
    doc,
    "这份讲稿和11页PPT对应。排练时不用逐字背，可以保留自己的说话习惯。每个人只要抓住本部分的开头、两三个重点和最后的交接句，整体时间就能控制在10分钟以内。",
    size=10.5,
    after=14,
    line=1.4,
)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(2.3), Cm(4.5), Cm(3.2), Cm(3.2)]
headers = ["主讲人", "负责内容", "对应页码", "建议时长"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=10, bold=True, color=BLACK)

rows = [
    ["1", "立项依据", "第1至4页", "约2分20秒"],
    ["2", "技术方案", "第5至6页", "约2分15秒"],
    ["3", "执行计划", "第7至8页", "约2分15秒"],
    ["4", "验收与保障", "第9至11页", "约3分钟"],
]
for ridx, row in enumerate(rows):
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].width = widths[i]
        set_cell_shading(cells[i], "FFFFFF" if ridx % 2 == 0 else PALE_GRAY)
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_run_font(r, size=10, color=BLACK)
set_table_borders(table)

add_paragraph(doc, "排练建议", style="Heading 1", size=14, bold=True, before=18, after=6, keep=True)
add_paragraph(doc, "第一次按原稿读，确认内容和页码；第二次只看每段第一句和关键词；第三次四个人连起来计时。交接时停半秒，看向下一位同学，再说交接句。", size=10.5, after=0)

add_page_break(doc)
add_speaker_header(doc, "一", "立项依据", "1至4", "约2分20秒")

add_cue(doc, "第1页  开场")
add_paragraph(doc, "各位老师、同学，大家好。我们组这次做的是任务二，科学投喂与鱼群摄食监测。这个题目听起来比较长，其实我们想解决的事情很具体，就是投喂的时候少靠感觉，多看鱼真实的反应。接下来我们四个人分别讲一下为什么要做、准备怎么做、十月份怎么安排，最后再说怎么验收。")

add_cue(doc, "第3页  为什么要做")
add_paragraph(doc, "先说一下现在的问题。网箱投喂很多时候还是靠现场人员看水花、看鱼群，再凭经验决定要不要继续投。投多了，饲料浪费，残饵还会沉到底部；投少了，鱼又吃不够，生长速度会受影响。")
add_paragraph(doc, "更麻烦的是，现场看起来很热闹，不一定代表鱼真的吃得好。水花大，可能只是鱼在表层活动；鱼聚在一起，也不一定说明投喂量正合适。再加上水下画面容易发绿、发暗，气泡和反光也会挡住目标，所以只看一个现象很容易判断错。")

add_cue(doc, "第4页  这次准备做到什么程度")
add_paragraph(doc, "所以我们把一次投喂分成几个环节来看。投喂前先看鱼群、水质和海况；投喂过程中看鱼游得快不快、聚得多不多，也结合水面和声呐信息；投喂结束后再看还有没有漂浮饵料或者底部残饵。最后，把系统建议、实际投了多少和投喂结果放到同一条记录里，之后能查、能比较。")
add_paragraph(doc, "不过这次只有一个月，我们不会一上来就做全自动投喂。十月份先把数据采集、画面处理、标签定义和小规模验证做通。这个基础打稳了，后面才有条件继续扩展。")

add_paragraph(doc, "我这部分先讲到这里。下面请第二位同学介绍我们具体准备怎么做。", size=11, bold=True, before=8, after=0)

add_page_break(doc)
add_speaker_header(doc, "二", "技术方案", "5至6", "约2分15秒")

add_cue(doc, "第5页  总体做法")
add_paragraph(doc, "我接着讲技术方案。我们的思路可以简单概括成一句话：多看几路信号，先把画面看清，再做判断。")
add_paragraph(doc, "最主要的数据还是水下视频，因为鱼的游动、聚集和抢食都能从视频里看到。如果设备有声呐，我们就用它补充鱼群密度和深度信息。水面影像可以看水花和饵料覆盖，水温、溶氧这些记录则用来解释为什么同样的投喂量，在不同环境下会有不同反应。")
add_paragraph(doc, "这些数据经过处理以后，我们会输出摄食状态的四个等级，再给出投喂量、速度和时长的建议。这里不会一开始就让系统直接控制设备，先把建议做出来，再和人工判断、实际投喂结果对照。")

add_cue(doc, "第6页  图像处理和数据规则")
add_paragraph(doc, "大家可以看这一页左边。原始水下画面比较绿，也比较模糊，鱼的轮廓和细节都不清楚。我们先用白平衡、对比度增强、去雾和去模糊这些传统方法做一个基线，再试轻量化模型。这样处理完以后，才能继续做鱼体识别、摄食判断和残饵检测。")
add_paragraph(doc, "数据这块有四条规则。第一，每次投喂都要有餐次编号，把时间、点位、网箱和实际投喂量对应起来。第二，摄食状态先用S0到S3四档表示，具体名称和边界第一周定下来。第三，训练集和测试集要按视频、网箱或者日期分开，不能把同一段视频的相邻画面拆到两边。第四，残饵区域由两个人分别标注，先看两个人标得是否一致，再评价模型。")
add_paragraph(doc, "我们会先做单一数据源的结果，再做多种数据融合。最后效果有没有提升，都按同一套数据和训练条件来比较。")

add_paragraph(doc, "技术方案就是这些。下面请第三位同学介绍十月份的安排和四个人怎么配合。", size=11, bold=True, before=8, after=0)

add_page_break(doc)
add_speaker_header(doc, "三", "执行计划", "7至8", "约2分15秒")

add_cue(doc, "第7页  十月安排")
add_paragraph(doc, "下面我来说一下具体怎么推进。项目时间就是整个十月，我们按四周来安排。")
add_paragraph(doc, "第一周先把任务说明、四档标签、餐次编号和参考方法定下来。这里如果定义不清楚，后面的数据越多，返工反而越多。")
add_paragraph(doc, "第二周主要练水下机器人，同时把拍摄规范定下来，比如分辨率、帧率、补光、拍摄距离和角度。第二周结束前如果拍摄还不稳定，我们就不急着进入正式采集。")
add_paragraph(doc, "第三周集中做数据采集和图像处理。现场条件允许就下水采集；条件不允许，就先用模拟环境，并用公开数据做补充。这个阶段要把原始数据整理好，也要保证清晰化程序能够完整跑通。")
add_paragraph(doc, "最后一周不再临时增加新功能，主要做指标计算、对比实验、错误分析和报告整理。这样能留出时间检查结果，而不是最后两天才开始拼材料。")

add_cue(doc, "第8页  四个人怎么配合")
add_paragraph(doc, "分工方面，第一位同学负责需求、标签和数据字段；第二位负责机器人操作、拍摄规范和原始数据；第三位负责图像增强、识别方法和运行脚本；第四位负责指标、错误分析和最后的报告整合。")
add_paragraph(doc, "每个人都要给下一位留下清楚的交接材料。比如采集前先确定要记录哪些字段，算法开始前先确认素材有没有按要求编号，评估前也要把结果和运行日志一起交过去。代码放在同一个仓库里，固定随机种子，每天提交一次。关键结果再由另外一位同学重新跑一遍。")

add_paragraph(doc, "执行安排介绍完了。下面请第四位同学讲一下我们准备怎么验收，以及遇到问题时怎么办。", size=11, bold=True, before=8, after=0)

add_page_break(doc)
add_speaker_header(doc, "四", "验收与保障", "9至11", "约3分钟")

add_cue(doc, "第9页  怎么验收")
add_paragraph(doc, "最后我来说验收和风险。我们这次不会只放几张效果比较好的图片，然后就说模型有效。最后能不能验收，主要看三类结果。")
add_paragraph(doc, "第一类是摄食状态识别。除了总体准确率，还要分别看四个等级的F1。分错一级和跨好几级，严重程度也不一样，所以这两种错误要分开统计。另外还要报告加权Kappa值和置信区间。")
add_paragraph(doc, "第二类是多种数据放在一起以后，到底有没有帮助。我们会把单独使用视频、声呐或者水面影像的结果，和融合后的结果放在一起比较。残饵识别除了看精确率和召回率，还要看面积误差，并把气泡、反光、饲料碎屑和鱼体遮挡造成的错误分开整理。")
add_paragraph(doc, "第三类是投喂建议能不能追溯。系统建议投多少，现场实际投了多少，最后有没有残饵，都要按餐次对应起来。水温或者溶氧条件不一样时，我们也会分开看表现。样本不够的情况就直接写样本不足，不勉强下结论。")
add_paragraph(doc, "这里有一个原则，某一项没达到预期并不等于整个项目没有价值，但必须把实际结果、样本量和原因说清楚，也要说明下一步准备怎么改。")

add_page_break(doc)
add_cue(doc, "第10页  最后交什么以及风险怎么处理")
add_paragraph(doc, "最终交付的东西分四类。第一类是任务书、四档定义和数据字典；第二类是运行环境、训练和推理脚本；第三类是完整结果、对比实验和错误样例；第四类是最后的项目报告，里面会写清楚哪些场景可以用，哪些场景目前还不能下结论。")
add_paragraph(doc, "风险方面，如果天气或者设备影响下水，我们就用模拟环境补上，并把公开数据单独标记。画面质量不好时，先调整距离、补光和防抖，不能把所有问题都推给后处理。如果声呐或水质数据缺失，就明确写出缺了什么，不做超出数据范围的判断。进度和版本方面，每天提交代码和记录，阶段结束做一次快照。")

add_cue(doc, "第11页  结尾")
add_paragraph(doc, "最后总结一下。我们希望用十月份把科学投喂最基本的数据链路跑通，范围不大，但每一步都能留下记录，也能让其他同学重新验证。我们申请按这个计划启动，同时希望协调水下机器人的使用时间、现场数据和一位独立复核人员。")
add_paragraph(doc, "我们的汇报到这里，谢谢大家。", size=11, bold=True, before=8, after=0)

doc.core_properties.title = "任务二科学投喂与鱼群摄食监测立项汇报讲稿"
doc.core_properties.subject = "四人汇报口语版讲稿"
doc.core_properties.author = "项目组"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(str(OUTPUT))
